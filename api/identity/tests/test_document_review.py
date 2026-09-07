"""
The stateless review endpoint.

Two things are load-bearing and both are asserted here: it reads the formats
other apps' members actually upload, and it never writes a row. The second is
the whole reason the endpoint exists — see review_views' module docstring —
so a regression that quietly starts persisting résumés has to fail a test.
"""

import io
import zipfile
from unittest.mock import MagicMock, patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from rest_framework.authtoken.models import Token

from identity.document_intake import DocumentUnreadable, extract_document_text
from identity.models import ProfessionalProfile, ResumeVersion, Skill

WELL_FORMED = """{
  "skills": [{"name": "Peer support", "category": "Community", "proficiency": "strong"}],
  "headline": "Community Health Worker",
  "email": "member@example.test",
  "phone": "502-555-0111",
  "linkedin_url": ""
}"""


def fake_response(text):
    block = MagicMock()
    block.type = "text"
    block.text = text
    response = MagicMock()
    response.content = [block]
    return response


def stub_stream(client, text):
    stream = client.return_value.messages.stream
    stream.return_value.__enter__.return_value.get_final_message.return_value = fake_response(text)
    return stream


def docx_bytes(paragraphs=("Community Health Worker", "Peer support, harm reduction")):
    from docx import Document

    document = Document()
    for line in paragraphs:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@override_settings(ANTHROPIC_API_KEY="test-key")
class DocumentReviewTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="service", password="x", email="service@example.test",
        )
        token, _ = Token.objects.get_or_create(user=self.user)
        self.client.defaults["HTTP_AUTHORIZATION"] = f"Token {token.key}"
        self.url = reverse("review-document")

    def post_file(self, name, data):
        return self.client.post(
            self.url, {"file": SimpleUploadedFile(name, data)}, format="multipart",
        )

    @patch("anthropic.Anthropic")
    def test_parses_a_docx_and_stores_nothing(self, client):
        stub_stream(client, WELL_FORMED)

        response = self.post_file("resume.docx", docx_bytes())

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            [s["name"] for s in response.json()["skills"]], ["Peer support"],
        )
        self.assertFalse(response.json()["stored"])
        # The point of the endpoint: no member data lands in this stack.
        self.assertEqual(ResumeVersion.objects.count(), 0)
        self.assertEqual(Skill.objects.count(), 0)
        self.assertEqual(ProfessionalProfile.objects.count(), 0)

    @patch("anthropic.Anthropic")
    def test_parses_plain_text_and_markdown(self, client):
        stub_stream(client, WELL_FORMED)
        body = b"Community Health Worker\n\nPeer support, harm reduction, intake\n" * 3

        for name in ("resume.txt", "resume.md"):
            with self.subTest(name=name):
                self.assertEqual(self.post_file(name, body).status_code, 200)

    @patch("anthropic.Anthropic")
    def test_accepts_already_extracted_text(self, client):
        stub_stream(client, WELL_FORMED)

        response = self.client.post(
            self.url,
            {"text": "Community Health Worker with peer support experience. " * 4},
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertGreater(response.json()["characters"], 50)

    def test_rejects_an_unsupported_type(self):
        response = self.post_file("payload.exe", b"MZ\x90\x00" * 20)

        self.assertEqual(response.status_code, 422)
        self.assertIn("PDF", response.json()["detail"])

    def test_rejects_a_renamed_file(self):
        response = self.post_file("not-really.pdf", b"MZ\x90\x00" * 20)

        self.assertEqual(response.status_code, 422)
        self.assertIn("real PDF", response.json()["detail"])

    def test_requires_a_file_or_text(self):
        response = self.client.post(self.url, {}, format="multipart")

        self.assertEqual(response.status_code, 400)
        self.assertIn(".docx", response.json()["accepted_extensions"])

    def test_rejects_an_anonymous_caller(self):
        self.client.defaults.pop("HTTP_AUTHORIZATION")

        self.assertEqual(self.post_file("resume.txt", b"hello " * 40).status_code, 401)


class DocumentIntakeTests(TestCase):
    def test_reads_docx_tables(self):
        from docx import Document

        document = Document()
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "2019-2024"
        table.rows[0].cells[1].text = "Outreach Coordinator"
        buffer = io.BytesIO()
        document.save(buffer)

        text = extract_document_text(buffer.getvalue(), "history.docx")

        self.assertIn("Outreach Coordinator", text)

    def test_strips_rtf_control_words(self):
        rtf = rb"{\rtf1\ansi\deff0 {\fonttbl{\f0 Helvetica;}}\f0\fs24 Peer support specialist}"

        text = extract_document_text(rtf, "resume.rtf")

        self.assertIn("Peer support specialist", text)
        self.assertNotIn("rtf1", text)

    def test_rejects_an_empty_file(self):
        with self.assertRaises(DocumentUnreadable):
            extract_document_text(b"", "resume.txt")

    def test_rejects_an_oversized_file(self):
        with self.assertRaises(DocumentUnreadable) as caught:
            extract_document_text(b"x" * (10 * 1024 * 1024 + 1), "resume.txt")

        self.assertIn("10MB", str(caught.exception))

    def test_rejects_a_docx_that_is_a_renamed_zip(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("something-else.txt", "hi")

        with self.assertRaises(DocumentUnreadable):
            extract_document_text(buffer.getvalue(), "resume.docx")
